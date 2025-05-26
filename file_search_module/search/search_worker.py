import re
import chardet
import docx # type: ignore
import openpyxl # type: ignore
from PyPDF2 import PdfReader # type: ignore
from PyQt6.QtCore import QObject, pyqtSignal
from pathlib import Path
# Assuming a logger might be used, e.g., from file_search_module.utils.logger import logger
# For now, using print for critical errors or relying on error_occurred signal.

class SearchWorker(QObject):
    progress_update = pyqtSignal(int)
    search_finished = pyqtSignal(str)
    error_occurred = pyqtSignal(str, str)
    file_matches_found = pyqtSignal(str, list)  # (file_path_str, [matching_lines])
    
    def __init__(self, folder_path_obj: Path, keyword: str): # Expects Path object
        super().__init__()
        self.folder_path_obj = folder_path_obj.resolve() # Ensure it's absolute
        self.keyword = keyword
        self.cancel_search_flag = False
        self.total_files = 0
        self.current_progress = 0
    
    def start_search(self):
        try:
            self.total_files = self._count_target_files(self.folder_path_obj)
        except Exception as e:
            self.error_occurred.emit("錯誤", f"計算檔案總數時發生錯誤: {e}")
            self.search_finished.emit("搜尋錯誤！")
            return

        if self.total_files == 0:
            self.search_finished.emit("找不到可搜尋的檔案！")
            return
        
        self.current_progress = 0 # Reset progress
        
        for item_path in self.folder_path_obj.rglob('*'):
            if self.cancel_search_flag:
                self.search_finished.emit("搜尋已取消！")
                return
            
            if item_path.is_file():
                ext = item_path.suffix.lower().lstrip('.')
                
                try:
                    if ext == "txt":
                        self._process_text_file(item_path, self.keyword)
                    elif ext == "pdf":
                        self._process_pdf_file(item_path, self.keyword)
                    elif ext == "docx":
                        self._process_docx_file(item_path, self.keyword)
                    elif ext == "xlsx":
                        self._process_xlsx_file(item_path, self.keyword)
                    elif ext in ("html", "htm"):
                        # Treat HTML/HTM as text files for keyword search
                        self._process_text_file(item_path, self.keyword)
                    # else:
                        # If not a target extension, we just increment progress for it if counted,
                        # or simply skip if not counted. Current logic counts only specific extensions.
                        # If we only want progress for matched extensions, this increment needs to be conditional.
                        # For simplicity, assuming all files are iterated for progress if counted in _count_target_files
                        # or that _count_target_files only counts relevant files.

                except Exception as e:
                    # Using str(item_path) for error messages
                    self.error_occurred.emit("處理錯誤", f"處理檔案 {str(item_path)} 時發生問題: {e}")
                
                # Increment progress for every file processed that was part of the count
                if ext in ("txt", "pdf", "docx", "xlsx", "html", "htm"):
                    self.current_progress += 1
                    self.progress_update.emit(self.current_progress)
        
        self.search_finished.emit("搜尋完成！")
    
    def cancel_search(self):
        self.cancel_search_flag = True
    
    def _count_target_files(self, folder_obj: Path) -> int: # folder_obj is Path
        count = 0
        for item_path in folder_obj.rglob('*'):
            if item_path.is_file():
                # Suffix includes the dot, e.g., ".txt"
                if item_path.suffix.lower().lstrip('.') in ("txt", "pdf", "docx", "xlsx", "html", "htm"):
                    count += 1
        return count
    
    def _process_text_file(self, file_path_obj: Path, keyword: str):
        try:
            with open(file_path_obj, "rb") as f: # open() handles Path objects
                raw_data = f.read()
            # Use chardet result, default to utf-8 if None or low confidence
            detected = chardet.detect(raw_data)
            detected_encoding = detected.get("encoding", "utf-8") if detected and detected['confidence'] and detected['confidence'] > 0.5 else "utf-8"
            if detected_encoding is None: detected_encoding = "utf-8" # Fallback

            text = raw_data.decode(detected_encoding, errors="replace") # Use 'replace' for robustness
        except Exception as e:
            self.error_occurred.emit("讀取錯誤", f"無法讀取文字檔 {str(file_path_obj)}: {e}")
            return
        
        lines = text.splitlines()
        matches = []
        for line_num, line_content in enumerate(lines): # line_num for future use if needed
            if re.search(re.escape(keyword), line_content, re.IGNORECASE):
                matches.append(line_content.strip())
        if matches:
            self.file_matches_found.emit(str(file_path_obj), matches)
    
    def _process_pdf_file(self, file_path_obj: Path, keyword: str):
        lines = []
        try:
            reader = PdfReader(file_path_obj) # PdfReader handles Path objects
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    lines.extend(text.splitlines()) # Use splitlines to be consistent
        except Exception as e:
            self.error_occurred.emit("PDF處理錯誤", f"無法處理 PDF {str(file_path_obj)}: {e}")
            return
        
        matches = []
        for line_content in lines:
            if re.search(re.escape(keyword), line_content, re.IGNORECASE):
                matches.append(line_content.strip())
        if matches:
            self.file_matches_found.emit(str(file_path_obj), matches)
    
    def _process_docx_file(self, file_path_obj: Path, keyword: str):
        try:
            doc = docx.Document(file_path_obj) # docx.Document handles Path objects
            lines = [para.text for para in doc.paragraphs]
        except Exception as e:
            self.error_occurred.emit("Word處理錯誤", f"無法處理 Word 文件 {str(file_path_obj)}: {e}")
            return
        
        matches = []
        for line_content in lines:
            if re.search(re.escape(keyword), line_content, re.IGNORECASE):
                matches.append(line_content.strip())
        if matches:
            self.file_matches_found.emit(str(file_path_obj), matches)
    
    def _process_xlsx_file(self, file_path_obj: Path, keyword: str):
        lines = []
        try:
            # openpyxl.load_workbook handles Path objects
            wb = openpyxl.load_workbook(file_path_obj, read_only=True, data_only=True) 
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows(values_only=True):
                    # Convert all cell values to string and join, then search
                    row_text = " ".join(str(cell) for cell in row if cell is not None)
                    lines.append(row_text)
        except Exception as e:
            self.error_occurred.emit("Excel處理錯誤", f"無法處理 Excel 文件 {str(file_path_obj)}: {e}")
            return
        
        matches = []
        for line_content in lines:
            if re.search(re.escape(keyword), line_content, re.IGNORECASE):
                matches.append(line_content.strip())
        if matches:
            self.file_matches_found.emit(str(file_path_obj), matches)
