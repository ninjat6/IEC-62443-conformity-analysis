import os
from pathlib import Path
from docx import Document as DocxDocument
from openpyxl import Workbook
from reportlab.pdfgen import canvas

# Define the directory for fixtures
# Assuming the script is run from the root of the repository or where 'tests' is a subdir
FIXTURES_DIR = Path("tests/fixtures")
FIXTURES_DIR.mkdir(parents=True, exist_ok=True)

# Create dummy.docx
docx_path = FIXTURES_DIR / "dummy.docx"
doc = DocxDocument()
doc.add_paragraph("Hello DOCX")
doc.save(docx_path)
print(f"Created: {docx_path}")

# Create empty.docx
empty_docx_path = FIXTURES_DIR / "empty.docx"
empty_doc = DocxDocument()
empty_doc.save(empty_docx_path)
print(f"Created: {empty_docx_path}")

# Create dummy.xlsx
xlsx_path = FIXTURES_DIR / "dummy.xlsx"
wb = Workbook()
sheet = wb.active
sheet["A1"] = "Hello XLSX"
wb.save(xlsx_path)
print(f"Created: {xlsx_path}")

# Create dummy.pdf
pdf_path = FIXTURES_DIR / "dummy.pdf"
c = canvas.Canvas(str(pdf_path))
c.drawString(72, 720, "Hello PDF") # Standard page size, draw near top
c.save()
print(f"Created: {pdf_path}")

# Create dummy.txt
txt_path = FIXTURES_DIR / "dummy.txt"
with open(txt_path, "w", encoding="utf-8") as f:
    f.write("Hello TXT")
print(f"Created: {txt_path}")

print("All dummy fixture files created.")
