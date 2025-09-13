# core/file_processor.py
from pathlib import Path
import docx
import openpyxl
import PyPDF2 # pdfplumber was mentioned in thought process but PyPDF2 is used in code
from conformity_analysis_module.utils.logger import logger
from typing import List

class FileProcessor:
    @staticmethod
    def extract_text_from_docx(file_path: Path) -> List[str]:
        """
        Extracts text snippets from a DOCX file.

        Args:
            file_path (Path): The path to the DOCX file.

        Returns:
            List[str]: A list of text snippets extracted from the document.
                       Returns an empty list if an error occurs or no text is found.
        """
        snippets: List[str] = []
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    snippets.append(text)
        except Exception as e:
            logger.error(f"處理 DOCX 檔案 {str(file_path)} 時發生錯誤: {e}")
        return snippets

    @staticmethod
    def extract_text_from_xlsx(file_path: Path) -> List[str]:
        """
        Extracts text snippets from an XLSX file.

        Args:
            file_path (Path): The path to the XLSX file.

        Returns:
            List[str]: A list of text snippets (typically cell contents joined per row)
                       extracted from the workbook. Returns an empty list if an error
                       occurs or no text is found.
        """
        snippets: List[str] = []
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join(str(cell) for cell in row if cell is not None).strip()
                    if row_text:
                        snippets.append(row_text)
        except Exception as e:
            logger.error(f"處理 XLSX 檔案 {str(file_path)} 時發生錯誤: {e}")
        return snippets

    @staticmethod
    def extract_text_from_pdf(file_path: Path) -> List[str]:
        """
        Extracts text snippets from a PDF file.

        Args:
            file_path (Path): The path to the PDF file.

        Returns:
            List[str]: A list of text snippets (lines of text) extracted from the PDF.
                       Returns an empty list if an error occurs or no text is found.
        """
        snippets: List[str] = []
        try:
            with open(file_path, "rb") as f:  # open() works with Path objects
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        for line in text.split('\n'):
                            line = line.strip()
                            if line:
                                snippets.append(line)
        except Exception as e:
            logger.error(f"處理 PDF 檔案 {str(file_path)} 時發生錯誤: {e}")
        return snippets

    @staticmethod
    def extract_text_snippets(file_path: Path) -> List[str]:
        """
        Extracts text snippets from a file based on its extension.

        Supports DOCX, XLSX, and PDF files. Logs a warning for unsupported file types.

        Args:
            file_path (Path): The path to the file.

        Returns:
            List[str]: A list of text snippets extracted from the file.
                       Returns an empty list if the file type is unsupported,
                       an error occurs during processing, or no text is found.
        """
        # Path.suffix includes the dot, e.g., ".docx"
        # We need to remove the dot for comparison.
        ext = file_path.suffix.lower()
        if ext: # Ensure suffix is not empty
            ext = ext[1:]

        if ext == 'docx':
            return FileProcessor.extract_text_from_docx(file_path)
        elif ext == 'xlsx':
            return FileProcessor.extract_text_from_xlsx(file_path)
        elif ext == 'pdf':
            return FileProcessor.extract_text_from_pdf(file_path)
        else:
            logger.warning(f"不支援的檔案類型: {str(file_path)} (副檔名: {ext})")
            return []