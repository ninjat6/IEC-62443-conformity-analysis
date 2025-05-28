# -*- coding: utf-8 -*-
# worksheet_updater.py
# -----------------------------------------------------------------------------
#  功能說明：
#  1. 讀取 analysis_results.json → KeywordAnalyzer 過濾。
#  2. 針對 **每一筆 snippet**：再次偵測引用檔案 (【編號 名稱】)。
#  3. 將結果寫入 IEC 62443-2-4 Worksheet：
#     • Conformity Statement
#     • Conformity Evidence：主文件 + "==>" 引用文件 (避免重複/自引用)。
#  4. 透過 PyQt6 讓使用者選擇輸出位置。
# -----------------------------------------------------------------------------

import json
import re
import shutil
import sys
from pathlib import Path
from PyQt6.QtWidgets import QApplication, QFileDialog

import openpyxl
from conformity_analysis_module.config import Config
from conformity_analysis_module.core.keyword_analyzer import KeywordAnalyzer
from conformity_analysis_module.utils.docx_section_extractor import DocxSectionExtractor
from conformity_analysis_module.utils.logger import logger


class WorksheetUpdater:
    """將分析結果寫入 IEC 62443-2-4 Excel 工作表。"""

    @staticmethod
    def update_worksheet():  # noqa: C901
        """執行 Worksheet 更新流程。"""

        # ------------------------------------------------------------------
        # 基本檔案檢查
        # ------------------------------------------------------------------
        analysis_file = Config.ANALYSIS_OUTPUT # Already a Path object from config refactor
        template_file = Config.WORKSHEET_FILE # Already a Path object from config refactor
        logger.info("analysis_results.json 路徑: %s", str(analysis_file))

        if not analysis_file.exists(): # Use Path.exists()
            return False, "analysis_results.json 不存在，請先執行分析"
        if not template_file.exists(): # Use Path.exists()
            return False, "找不到模板檔案 IEC62443_2_4d_2024-worksheet.xlsx"

        # ------------------------------------------------------------------
        # 讀取分析結果
        # ------------------------------------------------------------------
        try:
            with open(analysis_file, "r", encoding="utf-8") as fp:
                analysis_results = json.load(fp)
        except Exception:
            logger.exception("讀取 %s 時發生錯誤", analysis_file)
            return False, "讀取 analysis_results.json 時發生錯誤"

        # 建立全域 source_file 清單 (供引用偵測)
        all_source_files: list[str] = []
        for r in analysis_results:
            sf = r.get("source_file")
            if sf and sf not in all_source_files:
                all_source_files.append(sf)

        # ------------------------------------------------------------------
        # 關鍵詞過濾 & 每 snippet 引用偵測
        # ------------------------------------------------------------------
        analyzer = KeywordAnalyzer()
        enhanced = analyzer.analyze(analysis_results, min_score=2)

        ref_token = re.compile(r"【([^】]+)】")
        code_pat = re.compile(r"([A-Z]+-\d+-\d+(?:[A-Za-z])?)\s+.*")

        for res in enhanced:
            snippet = res.get("snippet", "")
            refs: list[str] = []
            for token in ref_token.findall(snippet):
                m = code_pat.match(token)
                if not m:
                    continue
                ref_code = m.group(1)
                base_code = re.sub(r"([A-Z]+-\d+-\d+)[A-Za-z]?", r"\1", ref_code)
                for sf_str in all_source_files: # sf_str is a string path
                    stem = Path(sf_str).stem # Use Path.stem
                    if ref_code in stem or base_code in stem:
                        refs.append(sf_str)
                        break
            res["referenced_files"] = refs

        # 統計
        total, accept = len(analysis_results), len(enhanced)
        reject = total - accept
        rej_rate = reject / total if total else 0
        logger.info(
            "關鍵詞分析: %d → 接受 %d, 拒絕 %d (%.2f%%)", total, accept, reject, rej_rate * 100
        )

        # ------------------------------------------------------------------
        # 複製模板
        # ------------------------------------------------------------------
        # Config.ROOT_DIR is already a Path object
        tmp_dir = Config.ROOT_DIR / "temp" 
        tmp_dir.mkdir(parents=True, exist_ok=True) # Use Path.mkdir
        tmp_file_path = tmp_dir / "IEC62443_2_4d_filled.xlsx" # Use / operator
        
        # shutil.copy works with Path objects
        shutil.copy(template_file, tmp_file_path) 

        # openpyxl.load_workbook works with Path objects
        wb = openpyxl.load_workbook(tmp_file_path) 
        ws = wb.active

        header = {
            str(ws.cell(row=1, column=c).value).strip().lower(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=1, column=c).value
        }
        id_col, stmt_col, evi_col = (
            header[h] for h in ("iec 62443-2-4 id", "conformity statement", "conformity evidence")
        )

        norm = lambda s: re.sub(r"\s+", "", str(s).upper())

        # 分組 by requirement
        by_req: dict[str, list[dict]] = {}
        for r in enhanced:
            req = norm(r.get("requirement", ""))
            by_req.setdefault(req, []).append(r)

        # ------------------------------------------------------------------
        # 寫入 Excel
        # ------------------------------------------------------------------
        for row in range(2, ws.max_row + 1):
            req_id = norm(ws.cell(row=row, column=id_col).value)
            if req_id not in by_req:
                continue

            entries = by_req[req_id]

            # ---------- Conformity Statement ----------
            stmt_cells: list[str] = []
            for ent in entries:
                source_file_path_obj = Path(ent["source_file"]) # ent["source_file"] is a string path
                src_name = source_file_path_obj.name # Use Path.name
                snippet = ent["snippet"]
                score = ent.get("keyword_score", 0)
                kw = ", ".join(ent.get("matched_keywords", []))
                section_no = ""
                if src_name.lower().endswith(".docx"): # src_name is already just the name string
                    try:
                        # DocxSectionExtractor expects a string path
                        section_no = DocxSectionExtractor(str(source_file_path_obj)).extract_sections(snippet)
                    except Exception:
                        section_no = "(無編號)"
                stmt_cells.append(
                    f"{src_name}\n Section {section_no} described that {snippet}\n[關鍵詞分數: {score}, 匹配關鍵詞: {kw}]"
                )
            ws.cell(row=row, column=stmt_col, value="\n\n".join(stmt_cells))

            # ---------- Conformity Evidence ----------
            evidence_texts: list[str] = []
            handled_main_docs: set[str] = set()  # 追蹤已處理的主文件
            all_ref_docs: set[str] = set()  # 追蹤所有被引用的文件
            doc_with_refs: dict[str, set[str]] = {}  # 每個主文件及其所有引用

            # 第一遍：收集所有主文件及其引用關係
            for ent in entries:
                source_file_path_obj = Path(ent["source_file"]) # ent["source_file"] is a string path
                src_name_stem = source_file_path_obj.stem # Use Path.stem for name without extension
                src_name_with_ext = source_file_path_obj.name # Use Path.name for full filename
                
                # 初始化這個文件的引用集合
                if src_name_stem not in doc_with_refs:
                    doc_with_refs[src_name_stem] = set()
                
                # 添加所有引用
                for rf_str in ent["referenced_files"]: # rf_str is a string path
                    ref_file_path_obj = Path(rf_str)
                    rf_name_stem = ref_file_path_obj.stem
                    
                    # 不添加自引用
                    if src_name_with_ext != ref_file_path_obj.name:
                        doc_with_refs[src_name_stem].add(rf_name_stem)
                        all_ref_docs.add(rf_name_stem)  # 標記為被引用文檔

            # 第二遍：構建結果文本
            # 先處理有snippet的主文件
            for ent in entries:
                source_file_path_obj = Path(ent["source_file"])
                src_name_stem = source_file_path_obj.stem
                
                # 如果這個主文件已經處理過，跳過
                if src_name_stem in handled_main_docs:
                    continue
                
                # 添加主文件及其引用
                text = src_name_stem
                if src_name_stem in doc_with_refs:
                    for ref_name_stem_sorted in sorted(doc_with_refs[src_name_stem]):
                        text += f"\n==> {ref_name_stem_sorted}"
                
                evidence_texts.append(text)
                handled_main_docs.add(src_name_stem)

            # 只添加那些有自己snippet但沒有被處理過的文檔
            for ent in entries:
                source_file_path_obj = Path(ent["source_file"])
                src_name_stem = source_file_path_obj.stem
                
                if src_name_stem not in handled_main_docs:
                    evidence_texts.append(src_name_stem)
                    handled_main_docs.add(src_name_stem)

            ws.cell(row=row, column=evi_col, value="\n\n".join(evidence_texts))

        # ------------------------------------------------------------------
        # 儲存 - 使用 PyQt6 文件對話框
        # ------------------------------------------------------------------
        
        # 確保 QApplication 實例存在
        app = QApplication.instance()
        if app is None:
            app = QApplication(sys.argv)
        
        # 使用 PyQt6 文件保存對話框
        save, _ = QFileDialog.getSaveFileName(
            None,  # 父窗口 (None = 獨立對話框)
            "選擇儲存 Excel 檔案",  # 對話框標題
            "IEC62443_2_4d_filled.xlsx",  # 預設檔名
            "Excel Files (*.xlsx);;All Files (*)"  # 文件類型過濾器
        )
        
        if not save:
            logger.warning("使用者取消存檔")
            return False, "使用者取消存檔"

        wb.save(save)
        logger.info("Worksheet 更新完成，已儲存至 %s", save)
        return (
            True,
            f"Worksheet 更新完成，檔案已儲存至 {save}\n"
            f"關鍵詞分析：接受 {accept} 筆, 拒絕 {reject} 筆 ({rej_rate * 100:.2f}%)",
        )